"""Generate two 13-slide, editable PowerPoint project demonstration decks."""
from pathlib import Path
import json
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
SLIDES = ROOT / 'slides'
SLIDES.mkdir(exist_ok=True)
DARK = RGBColor(17, 29, 48)
MID = RGBColor(35, 55, 77)
PAPER = RGBColor(250, 251, 253)
ACCENT = RGBColor(0, 134, 161)
AMBER = RGBColor(234, 139, 59)
TEXT = RGBColor(43, 55, 68)
MUTED = RGBColor(104, 117, 133)


def box(slide, x, y, w, h, fill, radius=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def txt(slide, content, x, y, w, h, size=17, color=TEXT, bold=False, align=None):
    tb=slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=tb.text_frame
    tf.clear()
    tf.word_wrap=True
    for i,line in enumerate(content.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line
        p.font.name='Aptos'
        p.font.size=Pt(size)
        p.font.bold=bold
        p.font.color.rgb=color
        p.space_after=Pt(9)
        if align is not None:p.alignment=align
    return tb


def base(prs, heading, number, titlebar=True):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    box(s, 0, 0, 13.33, 7.5, PAPER)
    box(s, 0, 0, .13, 7.5, ACCENT)
    txt(s, heading, .7, .45, 12.1, .72, 31, DARK, True)
    box(s, .7, 1.28, 11.9, .023, RGBColor(216,224,230))
    txt(s, 'MCA201P  •  COER University  •  Session 2026–27', .7, 7.08, 11.4, .25, 10, MUTED)
    txt(s, f'{number:02d} / 13', 11.8, 7.08, .75, .25, 10, MUTED)
    return s


def bullets(prs, heading, items, number, note=None):
    s=base(prs,heading,number)
    yy=1.68
    for item in items:
        box(s, .93, yy+.12, .09, .09, AMBER,True)
        txt(s,item,1.2,yy,11.1,.76,19)
        yy+=max(.85, min(1.17, .69 + len(item)/115))
    if note:
        box(s,.83,6.29,11.7,.55,RGBColor(232,242,245),True)
        txt(s,note,1.05,6.40,11.2,.34,12,DARK)
    return s


def cover(prs,subject,subtitle):
    s=prs.slides.add_slide(prs.slide_layouts[6])
    box(s,0,0,13.33,7.5,DARK)
    box(s,.73,1.02,.12,4.85,AMBER)
    txt(s,'MCA201P  |  LAB SHEET 10',1.18,1.18,10,.32,15,RGBColor(141,205,220),True)
    txt(s,subject,1.18,2,11.2,1.65,40,RGBColor(250,251,255),True)
    txt(s,subtitle,1.2,4.18,10.7,.8,20,RGBColor(200,218,231))
    txt(s,'Harsh Saini  •  253028036  •  COER University, Roorkee',1.2,6.55,11,.4,15,RGBColor(190,201,214))
    return s


def diagram(prs,heading,steps,number):
    s=base(prs,heading,number)
    for i,step in enumerate(steps):
        x=.62+i*2.55
        box(s,x,2.75,2.13,1.38,MID if i%2==0 else ACCENT,True)
        txt(s,step,x+.17,3.08,1.83,.9,15,RGBColor(255,255,255),True,PP_ALIGN.CENTER)
        if i<len(steps)-1:
            txt(s,'→',x+2.2,3.09,.32,.38,21,ACCENT,True,PP_ALIGN.CENTER)
    return s


def image_slide(prs, heading, img, number, caption):
    s=base(prs,heading,number)
    s.shapes.add_picture(str(img),Inches(1.08),Inches(1.52),width=Inches(11.1),height=Inches(4.95))
    txt(s,caption,1.0,6.55,11.5,.42,12,MUTED)
    return s


def metrics_slide(prs, heading, rows, number):
    s=base(prs,heading,number)
    y=1.75
    for label,value in rows:
        box(s,1.02,y,5.25,.74,RGBColor(232,239,244),True)
        txt(s,label,1.24,y+.2,4.8,.35,16,DARK,True)
        box(s,6.62,y,5.5,.74,RGBColor(243,247,249),True)
        txt(s,str(value),6.9,y+.18,4.8,.38,18,ACCENT,True)
        y+=.94
    return s


def build(task, filename):
    p=Presentation()
    p.slide_width=Inches(13.33);p.slide_height=Inches(7.5)
    loan = task=='loan'
    key='loan_approval' if loan else 'symptom_pattern'
    met=json.loads((ROOT/'outputs'/'lab_10'/f'{key}_metrics.json').read_text())
    title='Loan Approval Prediction' if loan else 'Symptom Pattern Classification'
    subtitle='Classification and interactive predictor' if loan else 'A simulated classification exercise — not a diagnosis'
    cover(p,title,subtitle)
    bullets(p,'Problem statement',[
        'Develop a trained classification system with complete data preprocessing.',
        'Compare different algorithms and report held-out evaluation metrics.',
        'Create a simple Streamlit prediction interface for practical demonstration.',
        'Important: this is an academic simulation, not a real-world decision tool.'
    ],2)
    bullets(p,'Project objectives',[
        'Prepare a reproducible source dataset with documented features.',
        'Prevent training/test leakage using scikit-learn pipelines.',
        'Compare Logistic Regression and Random Forest classifiers.',
        'Save models, metrics, confusion matrix and interactive user interface.'
    ],3)
    fields=['Monthly income, loan amount, credit history','Employment, education, marital status','Loan term and synthetic binary approval label'] if loan else [
        'Binary symptom indicators: fever, cough, sneezing and others',
        'Duration of symptoms in days',
        'Four fabricated condition-pattern labels; no clinical patients']
    bullets(p,'Dataset and features',fields+['Synthetic records: '+str(met['dataset_rows'])],4,
            'DATA LIMITATION: no real loan records or patient data.')
    diagram(p,'System architecture', ['CSV dataset','Clean and split','Fit pipelines','Evaluate & save','Streamlit UI'],5)
    bullets(p,'Data preprocessing and validation',[
        '80/20 stratified training and test split with fixed random state (42).',
        'Median imputation and standardization for numeric columns.',
        'Categorical columns: most-frequent imputation + one-hot encoding.',
        'All preprocessing is fitted only on the training split.'
    ],6)
    bullets(p,'Algorithms and comparison',[
        'Logistic Regression: linear discriminative classifier with regularization.',
        'Random Forest: nonlinear ensemble of decision trees.',
        'Compare accuracy, precision, recall and weighted F1.',
        'For binary loan classification, also report ROC-AUC.'
    ],7)
    metrics_slide(p,'Evaluation results — held-out test split',[
        ('Training rows',str(met['train_rows'])),('Test rows',str(met['test_rows'])),
        ('Selected model',met['best_model']),('Best weighted F1',f"{met['best_weighted_f1']:.3f}"),
        ('Data source','Synthetic / reproducible')],8)
    image_slide(p,'Confusion matrix — best model',ROOT/'outputs'/'lab_10'/f'{key}_confusion_matrix.png',9,
                'Counts describe synthetic test samples; they are not field validation.')
    bullets(p,'Prediction interface',[
        'Framework: Streamlit; starts with `streamlit run lab_10_app.py`.',
        'Validation: form inputs use constrained values and category choices.',
        'Uses saved trained preprocessing + classification pipeline.',
        'Displays predicted class with a prominent limitation notice.'
    ],10)
    bullets(p,'Observations and interpretation',[
        f"Best held-out weighted F1 in this simulated split: {met['best_weighted_f1']:.3f}.",
        f"Selected estimator: {met['best_model']} (based on weighted F1).",
        'Results are conditional on synthetic label-generation assumptions.',
        'No evidence here establishes real-world predictive validity.'
    ],11)
    bullets(p,'Risks, ethics and limitations',[
        'Synthetic data do not reflect operational prevalence or population bias.',
        'Never enter real financial details or medical data into this demo.',
        'Loan decisions require compliance, fairness audits and human governance.' if loan else
        'Medical diagnosis requires clinical validation and qualified medical professionals.',
        'Do not deploy as an automated high-stakes decision-maker.'
    ],12)
    bullets(p,'Conclusion and viva demonstration',[
        'Run `python lab_10.py` to regenerate both trained models.',
        'Open model-comparison CSV, JSON metrics and confusion matrix PNG.',
        'Launch `streamlit run lab_10_app.py`, select the project and predict.',
        'Explain pipeline, metrics, limitations and possible use of real validated data.'
    ],13)
    p.save(SLIDES/filename)
    return SLIDES/filename


if __name__ == '__main__':
    for t,n in [('loan','lab_10_loan_capstone_presentation.pptx'),
                ('symptoms','lab_10_symptom_capstone_presentation.pptx')]:
        print('Created',build(t,n))
